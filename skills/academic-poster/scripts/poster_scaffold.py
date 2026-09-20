#!/usr/bin/env python3
"""poster_scaffold.py — 從 poster-spec.json 產生可編輯的學術海報 .pptx。

用法:
    python poster_scaffold.py --spec poster-spec.json -o poster.pptx
    python poster_scaffold.py --verify --pptx poster.pptx --spec poster-spec.json

設計依據: ../references/poster-layout.md 的三層閱讀動線與尺寸速查表。
座標配置(three-column / better-poster 兩組)為本腳本依該文件版式骨架精神
自行推算的具體區塊座標(該文件僅有 ASCII 示意圖,未給精確百分比),詳見
skills/academic-poster/ATTRIBUTION.md 與 SKILL.md 的引用說明。

字型紀律(踩坑提醒,對應 harness LESSONS L-014 的同類坑):
    python-pptx 的 run.font.name 只會寫入 OOXML 的 <a:latin typeface="..."/>,
    不會同時寫入東亞字型 <a:ea typeface="..."/>。同一個 run 只要含有中文字元,
    就必須額外直接操作 run 的 rPr XML 補上 <a:ea typeface="標楷體"/>,否則中文
    在有正確東亞字型對應表的環境(如 PowerPoint)可能不會套用標楷體。
    本檔的 style_run() 一律：
      1) 設 run.font.name = "Times New Roman"(寫入 <a:latin>,英數字適用)
      2) 若 run.text 含 CJK 字元,額外補寫 <a:ea typeface="標楷體"/>
    不需要把中英文拆成不同 run——OOXML 本就是依字元類型(拉丁/東亞/複合文字)
    分別查三個 typeface 欄位來選字型,同一 run 內中英夾雜可以共用同一組 rPr。
"""
import argparse
import json
import os
import re
import sys

# L1 修復(2026-09-20)：包 try/except，stdout 被重導到不支援 reconfigure 的
# 物件時降級而非讓程式直接崩掉。
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

# ---------------------------------------------------------------------------
# 常數:尺寸表、字級門檻、字型、色彩
# ---------------------------------------------------------------------------

# (短邊cm, 長邊cm) —— 直式 = 短邊×長邊(寬×高);橫式互換。來源:references/poster-layout.md
SIZE_CM = {
    "A0": (84.1, 118.9),
    "A1": (59.4, 84.1),
}

CJK_RE = re.compile(r"[㐀-䶿一-鿿豈-﫿]")

LATIN_FONT = "Times New Roman"
EA_FONT = "標楷體"

# 字級門檻(pt),對應 SKILL.md Step2 三層閱讀動線表(A0 基準)。
# A1 版面較小,但 SKILL.md 未另訂 A1 專屬門檻,本腳本統一套用同一組門檻
# (寧可偏大也不違反"深讀層 ≥20pt 不可為塞內容而破"的紅線)。
MIN_TITLE = 85
MIN_KEY_FINDING = 60
MIN_SECTION_HEAD = 48
MIN_BODY = 28
MIN_DEEP = 24

# 實際套用字級(高於門檻留安全邊際)
FONT_TITLE = 90
FONT_KEY_FINDING = 66
FONT_SECTION_HEAD = 50
FONT_BODY = 30
FONT_DEEP = 26
FONT_CAPTION = 30  # 圖說用 body 級
FONT_MISC_MIN = 20  # 紅線:全海報字級絕不小於 20pt(logo 標籤等次要文字)
FONT_LOGO_LABEL = 20

# 色彩(1 主色 + 1 強調色 + 灰階,呼應 poster-layout.md 與 management-figure 的
# Okabe-Ito 色盲友善精神;此處僅取其中性深藍/赭紅兩色示意,非嚴格套用該色盤)
COLOR_PRIMARY = RGBColor(0x1F, 0x38, 0x64)      # 深藍:標題帶
COLOR_ACCENT = RGBColor(0xB3, 0x2B, 0x2B)       # 赭紅:主發現強調
COLOR_TEXT_DARK = RGBColor(0x22, 0x22, 0x22)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_BOX_BORDER = RGBColor(0x99, 0x99, 0x99)
COLOR_BOX_FILL = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_PLACEHOLDER_FILL = RGBColor(0xF0, 0xF0, 0xF0)

GAP = 0.012  # 版面內堆疊區塊間的縱向留白(以整頁高度的比例表示)


# ---------------------------------------------------------------------------
# 字型工具:同一 run 內視需要補上 <a:ea typeface="標楷體"/>
# ---------------------------------------------------------------------------

def style_run(run, size_pt, bold=False, color=None, align_ea=True):
    """設定 run 的字級/粗體/顏色,並確保中文能拿到標楷體(ea)、英數拿到 Times
    New Roman(latin)。踩坑提醒見檔頭說明——run.font.name 只寫 <a:latin>。
    """
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = LATIN_FONT  # 寫入 <a:latin typeface="Times New Roman"/>
    if color is not None:
        run.font.color.rgb = color
    if align_ea and run.text and CJK_RE.search(run.text):
        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = OxmlElement("a:ea")
            rPr.append(ea)
        ea.set("typeface", EA_FONT)
    return run


def add_paragraph_text(text_frame, text, size_pt, bold=False, color=None,
                        align=PP_ALIGN.LEFT, first=False):
    """在 text_frame 加一段文字(單一 run),套用字型紀律。"""
    p = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
    p.alignment = align
    run = p.add_run()
    run.text = text
    style_run(run, size_pt, bold=bold, color=color)
    return p


# ---------------------------------------------------------------------------
# 版面工具:以「整頁比例」描述座標,再依實際海報尺寸換算成 EMU
# ---------------------------------------------------------------------------

def frac_rect_to_emu(rect, slide_w, slide_h):
    x, y, w, h = rect
    return (Emu(int(x * slide_w)), Emu(int(y * slide_h)),
            Emu(int(w * slide_w)), Emu(int(h * slide_h)))


def stack(x, y0, y1, w, weights, gap=GAP):
    """把 [y0, y1] 區間依 weights 比例切成多個等寬堆疊區塊(含區塊間留白 gap)。
    回傳一組 (x, y, w, h) fraction rect。"""
    n = len(weights)
    total_gap = gap * (n - 1) if n > 1 else 0.0
    avail = max((y1 - y0) - total_gap, 0.0)
    total_w = sum(weights)
    rects = []
    cur = y0
    for wt in weights:
        h = avail * (wt / total_w)
        rects.append((x, cur, w, h))
        cur += h + gap
    return rects


# ---------------------------------------------------------------------------
# 兩組版式的區塊座標配置(依 references/poster-layout.md 的 ASCII 版面骨架
# 精神推算為具體 fraction 座標;該文件只有示意圖沒有精確數字,此為本腳本的
# 合理推算,已在 ATTRIBUTION.md / SKILL.md 註明)
# ---------------------------------------------------------------------------

def build_blocks_three_column():
    """傳統三欄式:標題帶→主發現帶→三欄(背景/假說 | 架構圖/方法 | 結果圖/結果/
    結論/QR)→底欄(參考文獻)。"""
    M = 0.03          # 頁邊距
    GAP_COL = 0.02    # 欄間距
    col_w = (1.0 - 2 * M - 2 * GAP_COL) / 3
    x0, x1, x2 = M, M + col_w + GAP_COL, M + 2 * (col_w + GAP_COL)

    blocks = []
    blocks.append({"kind": "title", "rect": (M, 0.01, (1 - 2 * M) * 0.78, 0.075)})
    blocks.append({"kind": "logos", "rect": (M + (1 - 2 * M) * 0.80, 0.01, (1 - 2 * M) * 0.20, 0.075)})
    blocks.append({"kind": "key_finding", "rect": (M, 0.095, 1 - 2 * M, 0.06)})

    col_top, col_bottom = 0.17, 0.925

    r = stack(x0, col_top, col_bottom, col_w, [1, 1])
    blocks.append({"kind": "section", "rect": r[0], "section_id": "background"})
    blocks.append({"kind": "section", "rect": r[1], "section_id": "hypotheses"})

    r = stack(x1, col_top, col_bottom, col_w, [0.85, 1.15])
    blocks.append({"kind": "figure", "rect": r[0], "figure_id": "framework"})
    blocks.append({"kind": "section", "rect": r[1], "section_id": "methods"})

    r = stack(x2, col_top, col_bottom, col_w, [0.8, 0.65, 0.65, 0.9])
    blocks.append({"kind": "figure", "rect": r[0], "figure_id": "results_fig"})
    blocks.append({"kind": "section", "rect": r[1], "section_id": "results"})
    blocks.append({"kind": "section", "rect": r[2], "section_id": "discussion"})
    blocks.append({"kind": "qr_contact", "rect": r[3]})

    blocks.append({"kind": "footer_refs", "rect": (M, 0.935, 1 - 2 * M, 0.055)})
    return blocks


def build_blocks_better_poster():
    """Better Poster 式:薄標題帶→左側欄(背景/假說/小架構圖) + 中央大字主發現與
    主圖 + 右側欄(方法/結果數字卡/結論) + 大 QR → 底欄(參考文獻)。"""
    M = 0.02
    blocks = []
    blocks.append({"kind": "title", "rect": (M, 0.01, 0.70, 0.045)})
    blocks.append({"kind": "logos", "rect": (0.74, 0.01, 1 - M - 0.74, 0.045)})

    left_x, left_w = M, 0.23
    r = stack(left_x, 0.07, 0.93, left_w, [0.9, 0.9, 1.2])
    blocks.append({"kind": "section", "rect": r[0], "section_id": "background"})
    blocks.append({"kind": "section", "rect": r[1], "section_id": "hypotheses"})
    blocks.append({"kind": "figure", "rect": r[2], "figure_id": "framework"})

    center_x, center_w = 0.27, 0.46
    r = stack(center_x, 0.07, 0.93, center_w, [0.32, 0.68])
    blocks.append({"kind": "key_finding", "rect": r[0], "big": True})
    blocks.append({"kind": "figure", "rect": r[1], "figure_id": "results_fig", "big": True})

    right_x, right_w = 0.75, 1 - M - 0.75
    r = stack(right_x, 0.07, 0.80, right_w, [1, 1, 1])
    blocks.append({"kind": "section", "rect": r[0], "section_id": "methods"})
    blocks.append({"kind": "section", "rect": r[1], "section_id": "results"})
    blocks.append({"kind": "section", "rect": r[2], "section_id": "discussion"})

    blocks.append({"kind": "qr_contact", "rect": (right_x, 0.81, right_w, 0.12), "big_qr": True})
    blocks.append({"kind": "footer_refs", "rect": (M, 0.935, 1 - 2 * M, 0.055)})
    return blocks


LAYOUT_BUILDERS = {
    "three-column": build_blocks_three_column,
    "better-poster": build_blocks_better_poster,
}


# ---------------------------------------------------------------------------
# 繪製
# ---------------------------------------------------------------------------

def add_box(slide, rect_emu, fill=None, line_color=COLOR_BOX_BORDER, dashed=False):
    left, top, width, height = rect_emu
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
        if dashed:
            ln = shape.line._get_or_add_ln()
            prstDash = OxmlElement("a:prstDash")
            prstDash.set("val", "dash")
            ln.append(prstDash)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Cm(0.4)
    tf.margin_right = Cm(0.4)
    tf.margin_top = Cm(0.3)
    tf.margin_bottom = Cm(0.3)
    return shape


def render_title(slide, rect_emu, spec, name):
    shape = add_box(slide, rect_emu, fill=COLOR_PRIMARY, line_color=None)
    shape.name = name
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_paragraph_text(tf, spec["title"], FONT_TITLE, bold=True,
                        color=COLOR_WHITE, align=PP_ALIGN.LEFT, first=True)
    authors = spec.get("authors", "")
    institution = spec.get("institution", "")
    if authors or institution:
        sub = " ・ ".join([s for s in (authors, institution) if s])
        add_paragraph_text(tf, sub, FONT_DEEP, bold=False, color=COLOR_WHITE)


def render_key_finding(slide, rect_emu, spec, name, big=False):
    shape = add_box(slide, rect_emu, fill=COLOR_ACCENT, line_color=None)
    shape.name = name
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    size = FONT_KEY_FINDING + (20 if big else 0)
    add_paragraph_text(tf, spec["key_finding"], size, bold=True,
                        color=COLOR_WHITE, align=PP_ALIGN.CENTER, first=True)


def render_section(slide, rect_emu, section, name):
    shape = add_box(slide, rect_emu, fill=COLOR_BOX_FILL, line_color=COLOR_BOX_BORDER)
    shape.name = name
    tf = shape.text_frame
    add_paragraph_text(tf, section["heading"], FONT_SECTION_HEAD, bold=True,
                        color=COLOR_PRIMARY, align=PP_ALIGN.LEFT, first=True)
    add_paragraph_text(tf, section["body"], FONT_BODY, bold=False,
                        color=COLOR_TEXT_DARK, align=PP_ALIGN.LEFT)


def render_figure(slide, rect_emu, figure, name, spec_dir):
    left, top, width, height = rect_emu
    image_path = figure.get("image_path") or ""
    resolved = os.path.join(spec_dir, image_path) if image_path else ""
    if image_path and os.path.isfile(resolved):
        pic_height = Emu(int(height * 0.82))
        cap_height = Emu(int(height * 0.18))
        pic = slide.shapes.add_picture(resolved, left, top, width=width, height=pic_height)
        pic.name = name
        cap_top = Emu(int(top + pic_height))
        cap_box = add_box(slide, (left, cap_top, width, cap_height),
                           fill=None, line_color=None)
        cap_box.name = name + ":caption"
        add_paragraph_text(cap_box.text_frame, figure["caption"], FONT_CAPTION,
                            color=COLOR_TEXT_DARK, align=PP_ALIGN.CENTER, first=True)
        return
    # 找不到圖檔:畫佔位框,不假裝有圖、不靜默略過(誠實揭露缺漏)
    shape = add_box(slide, rect_emu, fill=COLOR_PLACEHOLDER_FILL,
                     line_color=COLOR_BOX_BORDER, dashed=True)
    shape.name = name
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    label = "[圖檔佔位:{}]".format(image_path) if image_path else "[圖檔佔位]"
    add_paragraph_text(tf, label, FONT_BODY, bold=True,
                        color=COLOR_BOX_BORDER, align=PP_ALIGN.CENTER, first=True)
    add_paragraph_text(tf, figure["caption"], FONT_BODY, color=COLOR_TEXT_DARK,
                        align=PP_ALIGN.CENTER)


def render_qr_contact(slide, rect_emu, spec, name, big_qr=False):
    shape = add_box(slide, rect_emu, fill=COLOR_BOX_FILL, line_color=COLOR_BOX_BORDER)
    shape.name = name
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    label = "QR CODE(佔位,未產生實際條碼——見腳本已知限制)" if big_qr else "QR CODE(佔位)"
    add_paragraph_text(tf, label, FONT_DEEP, bold=True,
                        color=COLOR_PRIMARY, align=PP_ALIGN.CENTER, first=True)
    add_paragraph_text(tf, spec.get("qr_url", ""), FONT_DEEP,
                        color=COLOR_TEXT_DARK, align=PP_ALIGN.CENTER)
    add_paragraph_text(tf, spec.get("contact", ""), FONT_DEEP,
                        color=COLOR_TEXT_DARK, align=PP_ALIGN.CENTER)


def render_footer_refs(slide, rect_emu, spec, name):
    shape = add_box(slide, rect_emu, fill=None, line_color=COLOR_BOX_BORDER)
    shape.name = name
    tf = shape.text_frame
    ref_section = next((s for s in spec["sections"] if s.get("id") == "references"), None)
    text = ref_section["body"] if ref_section else ""
    add_paragraph_text(tf, text, FONT_DEEP, color=COLOR_TEXT_DARK,
                        align=PP_ALIGN.LEFT, first=True)


def render_logos(slide, rect_emu, spec, name, spec_dir):
    left, top, width, height = rect_emu
    logos = spec.get("logos", [])
    if not logos:
        return
    n = len(logos)
    slot_w = Emu(int(width / n))
    for i, logo_path in enumerate(logos):
        slot_left = Emu(int(left + i * slot_w))
        resolved = os.path.join(spec_dir, logo_path)
        if os.path.isfile(resolved):
            slide.shapes.add_picture(resolved, slot_left, top, width=slot_w, height=height)
        else:
            shape = add_box(slide, (slot_left, top, slot_w, height),
                             fill=COLOR_PLACEHOLDER_FILL, line_color=COLOR_BOX_BORDER,
                             dashed=True)
            shape.name = "{}:{}".format(name, i)
            tf = shape.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            base = os.path.basename(logo_path)
            add_paragraph_text(tf, "[logo:{}]".format(base), FONT_LOGO_LABEL,
                                color=COLOR_BOX_BORDER, align=PP_ALIGN.CENTER, first=True)


# ---------------------------------------------------------------------------
# 主建置流程
# ---------------------------------------------------------------------------

def get_slide_dimensions_cm(size, orientation):
    if size not in SIZE_CM:
        raise ValueError("不支援的 size: {}(僅支援 {}）".format(size, list(SIZE_CM.keys())))
    short_side, long_side = SIZE_CM[size]
    if orientation == "portrait":
        return short_side, long_side
    elif orientation == "landscape":
        return long_side, short_side
    raise ValueError("不支援的 orientation: {}(僅支援 portrait/landscape）".format(orientation))


def build_poster(spec, spec_dir):
    layout = spec.get("layout")
    if layout not in LAYOUT_BUILDERS:
        raise ValueError("不支援的 layout: {}(僅支援 {}）".format(layout, list(LAYOUT_BUILDERS.keys())))

    width_cm, height_cm = get_slide_dimensions_cm(spec["size"], spec["orientation"])

    prs = Presentation()
    prs.slide_width = Cm(width_cm)
    prs.slide_height = Cm(height_cm)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版面

    slide_w, slide_h = prs.slide_width, prs.slide_height
    sections_by_id = {s["id"]: s for s in spec["sections"]}
    figures_by_id = {f["id"]: f for f in spec["figures"]}

    blocks = LAYOUT_BUILDERS[layout]()
    placed_section_ids, placed_figure_ids = set(), set()

    for block in blocks:
        rect_emu = frac_rect_to_emu(block["rect"], slide_w, slide_h)
        kind = block["kind"]
        if kind == "title":
            render_title(slide, rect_emu, spec, "poster:title")
        elif kind == "key_finding":
            render_key_finding(slide, rect_emu, spec, "poster:key_finding",
                                big=block.get("big", False))
        elif kind == "section":
            sid = block["section_id"]
            section = sections_by_id.get(sid)
            if section is None:
                print("[警告] spec 缺少 section id={}，該區塊留空".format(sid), file=sys.stderr)
                continue
            render_section(slide, rect_emu, section, "poster:section:{}".format(sid))
            placed_section_ids.add(sid)
        elif kind == "figure":
            fid = block["figure_id"]
            figure = figures_by_id.get(fid)
            if figure is None:
                print("[警告] spec 缺少 figure id={}，該區塊留空".format(fid), file=sys.stderr)
                continue
            render_figure(slide, rect_emu, figure, "poster:figure:{}".format(fid), spec_dir)
            placed_figure_ids.add(fid)
        elif kind == "qr_contact":
            render_qr_contact(slide, rect_emu, spec, "poster:qr_contact",
                               big_qr=block.get("big_qr", False))
        elif kind == "footer_refs":
            render_footer_refs(slide, rect_emu, spec, "poster:footer_refs")
        elif kind == "logos":
            render_logos(slide, rect_emu, spec, "poster:logos", spec_dir)
        else:
            raise ValueError("未知的 block kind: {}".format(kind))

    # 不靜默丟資料:spec 裡有但版式沒放到的 section/figure 明確回報
    all_section_ids = set(sections_by_id.keys()) - {"references"}  # references 走 footer_refs
    missing_sections = all_section_ids - placed_section_ids
    missing_figures = set(figures_by_id.keys()) - placed_figure_ids
    if missing_sections:
        print("[警告] 版式 {} 未使用到的 sections: {}".format(layout, sorted(missing_sections)),
              file=sys.stderr)
    if missing_figures:
        print("[警告] 版式 {} 未使用到的 figures: {}".format(layout, sorted(missing_figures)),
              file=sys.stderr)

    return prs


# ---------------------------------------------------------------------------
# 讀回驗證(--verify)
# ---------------------------------------------------------------------------

CATEGORY_THRESHOLDS = {
    "title": MIN_TITLE,
    "key_finding": MIN_KEY_FINDING,
    "section_head": MIN_SECTION_HEAD,
    "body": MIN_BODY,
    "deep": MIN_DEEP,
}


def iter_runs_by_shape(slide):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        yield shape, shape.text_frame


def verify(pptx_path, spec):
    errors = []
    findings = {k: [] for k in CATEGORY_THRESHOLDS}

    prs = Presentation(pptx_path)
    expect_w_cm, expect_h_cm = get_slide_dimensions_cm(spec["size"], spec["orientation"])
    actual_w_cm = prs.slide_width / 360000
    actual_h_cm = prs.slide_height / 360000

    print("頁面尺寸讀回: {:.2f} cm x {:.2f} cm (預期 {:.2f} x {:.2f})".format(
        actual_w_cm, actual_h_cm, expect_w_cm, expect_h_cm))
    if abs(actual_w_cm - expect_w_cm) > 0.05 or abs(actual_h_cm - expect_h_cm) > 0.05:
        errors.append("頁面尺寸不符: 讀回 {:.3f}x{:.3f} cm，預期 {:.3f}x{:.3f} cm".format(
            actual_w_cm, actual_h_cm, expect_w_cm, expect_h_cm))

    assert len(prs.slides) == 1, "應只有單頁海報"
    slide = prs.slides[0]

    ea_evidence = None
    for shape, tf in iter_runs_by_shape(slide):
        name = shape.name or ""
        for para_idx, para in enumerate(tf.paragraphs):
            for run in para.runs:
                if run.font.size is None:
                    continue
                size_pt = run.font.size.pt
                if name.startswith("poster:title"):
                    findings["title"].append((name, size_pt))
                elif name.startswith("poster:key_finding"):
                    findings["key_finding"].append((name, size_pt))
                elif name.startswith("poster:section:"):
                    # 區塊標題(第一段)算 section_head，其餘段落算 body
                    is_heading = (para_idx == 0)
                    findings["section_head" if is_heading else "body"].append((name, size_pt))
                elif name.startswith("poster:figure:"):
                    findings["body"].append((name, size_pt))
                elif name.startswith("poster:qr_contact") or name.startswith("poster:footer_refs"):
                    findings["deep"].append((name, size_pt))
                if ea_evidence is None and CJK_RE.search(run.text or ""):
                    rPr = run._r.find(qn("a:rPr"))
                    if rPr is not None:
                        ea = rPr.find(qn("a:ea"))
                        if ea is not None:
                            from lxml import etree
                            ea_evidence = (name, run.text[:20], etree.tostring(rPr).decode("utf-8"))

    print()
    for cat, threshold in CATEGORY_THRESHOLDS.items():
        sizes = [s for _, s in findings[cat]]
        if not sizes:
            errors.append("類別 {} 找不到任何文字框可供抽查".format(cat))
            print("[{}] 無樣本".format(cat))
            continue
        max_size = max(sizes)
        sample_name = [n for n, s in findings[cat] if s == max_size][0]
        print("[{}] 門檻>={}pt，抽查到的字級樣本(前5筆)={}，最大值={}pt (來自 {})".format(
            cat, threshold, sizes[:5], max_size, sample_name))
        if max_size < threshold:
            errors.append("類別 {} 最大字級 {}pt 未達門檻 {}pt".format(cat, max_size, threshold))

    print()
    if ea_evidence:
        name, text_sample, xml = ea_evidence
        print("ea 字型設定證據 (shape={}, text='{}'):\n{}".format(name, text_sample, xml))
    else:
        errors.append("找不到任何中文 run 帶有 <a:ea> 字型設定")

    if errors:
        print("\n驗證失敗:")
        for e in errors:
            print(" - " + e)
        raise SystemExit(1)
    print("\n驗證通過:所有斷言成立。")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="從 poster-spec.json 產生/驗證學術海報 .pptx")
    parser.add_argument("--spec", required=True, help="poster-spec.json 路徑")
    parser.add_argument("-o", "--output", default=None,
                         help="輸出 .pptx 路徑(預設 ./output/poster.pptx，M4 修復：不直接寫 cwd)")
    parser.add_argument("--verify", action="store_true", help="改為驗證模式:讀回 --output 指定的 pptx")
    parser.add_argument("--pptx", help="--verify 模式下要讀回檢查的 pptx 路徑(預設沿用 -o)")
    args = parser.parse_args()

    if args.output is None:
        os.makedirs("output", exist_ok=True)
        args.output = os.path.join("output", "poster.pptx")
        print(f"[提醒] 未指定 -o/--output，輸出預設寫到 {args.output}")

    spec_path = os.path.abspath(args.spec)
    spec_dir = os.path.dirname(spec_path)
    with open(spec_path, "r", encoding="utf-8") as f:
        spec = json.load(f)

    if args.verify:
        pptx_path = args.pptx or args.output
        # 單指令 `--spec X -o Y --verify` 應「先建再驗」；只有明確給 --pptx 指向既有檔才跳過建檔
        # （驗收 B8：舊版直接跳過 build，單指令會因檔案不存在而 PackageNotFoundError）
        if not args.pptx or not os.path.exists(pptx_path):
            prs = build_poster(spec, spec_dir)
            prs.save(pptx_path)
            print("已產生: {} (layout={}, size={}, orientation={})".format(
                pptx_path, spec["layout"], spec["size"], spec["orientation"]))
        verify(pptx_path, spec)
        return

    prs = build_poster(spec, spec_dir)
    prs.save(args.output)
    print("已產生: {} (layout={}, size={}, orientation={})".format(
        args.output, spec["layout"], spec["size"], spec["orientation"]))


if __name__ == "__main__":
    main()
