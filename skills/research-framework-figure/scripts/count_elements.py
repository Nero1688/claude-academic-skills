#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
count_elements.py — SVG／PPTX 元素清點工具（回歸與驗收用）

【為什麼需要這支】
research-framework-figure 的 PPTX 匯出曾經靜默丟掉全部箭頭與假說標籤——SVG 有
完整的線與字，PPTX 卻只剩方框。當時的煙霧測試只驗證「產檔不報錯」，沒有驗證
內容是否正確，因此沒抓到這個缺陷。本工具改採「逐格式清點元素數」：把 SVG 與
PPTX 的圖形元素都數過、分類，讓回歸測試（改版前後元素數必須相同）與新版式驗收
（箭頭與文字標籤兩邊都要有）可以量化比對，而不是憑「腳本沒報錯」心證。

【SVG 清點】
計數 rect / path / line / text；「箭頭類」的判準是**有沒有 marker-end**
（真的有箭頭頭），不是元素種類——path 與 line 都可能是箭頭類或裝飾類：
  - path：有 marker-end 者為箭頭路徑，其餘（例如箭頭標記定義本身）另計。
  - line：有 marker-end 者為箭頭，其餘（框標題分隔線、直式階段標籤引導
    線等）歸「裝飾線」，不計入箭頭類小計（此前版本誤把所有 line 都算進
    箭頭類，會虛灌數字、與 PPTX 端對不起來，已修正）。
  - text 再分：class="hypo" 的假說標籤 vs 一般文字。

【PPTX 清點】
用 python-pptx 讀回所有 slide 的 shapes，依 shape 名稱與型別分類：
  - connector/freeform（箭頭類）：shape.name 以 "conn" 開頭（本技能自訂的
    svg2drawingml 轉換命名慣例），或 shape_type 為 FREEFORM。
  - textbox：shape_type 為 TEXT_BOX（例如假說標籤、直式階段標籤）。
  - autoshape 含文字／無文字：其餘的方框（例如研究架構圖的變數框），依
    text_frame 是否有非空文字區分。
「文字類」小計 = textbox + autoshape 含文字（凡是承載得到文字的形狀都算）。

【用法】
    python count_elements.py fig.svg
    python count_elements.py fig.pptx
    python count_elements.py fig.svg fig.pptx --expect-svg 42 --expect-pptx 18

給了 --expect-svg / --expect-pptx，清點後的總數若不等於期望值就印出差異並以
非零 exit code 結束（供 CI／驗收腳本串接判斷成功與否，不必肉眼比對）。

last_verified: 2026-09-20
"""
from __future__ import annotations

import argparse
import sys
import xml.etree.ElementTree as ET

# Windows 主控台需要這行才能正確印中文；L1 修復(2026-09-20)：包 try/except，
# stdout 被重導到不支援 reconfigure 的物件時降級而非讓程式直接崩掉。
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

SVG_NS = "{http://www.w3.org/2000/svg}"


def count_svg(path: str) -> dict:
    """清點一個 SVG 檔的 rect / path / line / text，並依屬性再細分。

    「箭頭類」的判準是**有沒有 `marker-end`（真的有箭頭頭）**，不是「是不是
    line」。本技能的 `arrow()` 一律用 `<path class="conn" marker-end="...">`
    畫箭頭；`<line>` 只用在裝飾性場合（框標題與內容的分隔線、sample_flow
    左側直式階段標籤的引導線），從來不是箭頭。舊版誤把「所有 line」都算進
    箭頭類小計，會虛灌 SVG 側的箭頭數、讓它跟 PPTX 的 connector/freeform
    數（只承接真箭頭）對不起來——這裡改成依 marker-end 屬性判斷，並把沒有
    箭頭頭的線另外歸一類（裝飾線），不再混進箭頭類。
    """
    tree = ET.parse(path)
    root = tree.getroot()

    c = {
        "rect": 0,
        "path": 0, "path_arrow": 0, "path_other": 0,
        "line": 0, "line_arrow": 0, "line_decorative": 0,
        "text": 0, "text_hypo": 0, "text_plain": 0,
    }
    for el in root.iter():
        tag = el.tag.replace(SVG_NS, "")
        cls = el.get("class", "")
        has_arrowhead = bool(el.get("marker-end"))
        if tag == "rect":
            c["rect"] += 1
        elif tag == "path":
            c["path"] += 1
            if has_arrowhead:
                c["path_arrow"] += 1
            else:
                c["path_other"] += 1
        elif tag == "line":
            c["line"] += 1
            if has_arrowhead:
                c["line_arrow"] += 1
            else:
                c["line_decorative"] += 1
        elif tag == "text":
            c["text"] += 1
            if cls == "hypo":
                c["text_hypo"] += 1
            else:
                c["text_plain"] += 1

    c["arrow_like_total"] = c["path_arrow"] + c["line_arrow"]  # 只算真的有箭頭頭的
    c["decorative_line_total"] = c["line_decorative"]
    c["text_total"] = c["text"]
    c["total"] = c["rect"] + c["path"] + c["line"] + c["text"]
    return c


def count_pptx(path: str) -> dict:
    """清點一個 PPTX 檔所有 slide 的 shapes，依名稱與型別分類。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    c = {
        "connector_freeform": 0,
        "textbox": 0,
        "autoshape_with_text": 0,
        "autoshape_no_text": 0,
        "other": 0,
        "shapes_total": 0,
        "paragraphs_total": 0,
    }
    for slide in prs.slides:
        for shp in slide.shapes:
            c["shapes_total"] += 1
            name = shp.name or ""

            has_text = False
            try:
                has_text = bool(shp.has_text_frame) and shp.text_frame.text.strip() != ""
            except Exception:
                has_text = False
            if has_text:
                c["paragraphs_total"] += len(shp.text_frame.paragraphs)

            try:
                shp_type = shp.shape_type
            except (NotImplementedError, ValueError, KeyError):
                shp_type = None

            is_connector = name.startswith("conn") or shp_type == MSO_SHAPE_TYPE.FREEFORM
            if is_connector:
                c["connector_freeform"] += 1
            elif shp_type == MSO_SHAPE_TYPE.TEXT_BOX:
                c["textbox"] += 1
            elif has_text:
                c["autoshape_with_text"] += 1
            elif shp_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                c["autoshape_no_text"] += 1
            else:
                c["other"] += 1

    c["arrow_like_total"] = c["connector_freeform"]
    c["text_total"] = c["textbox"] + c["autoshape_with_text"]
    return c


def _print_svg_report(f: str, c: dict) -> None:
    print(f"\n[SVG] {f}")
    print(f"  rect                        : {c['rect']}")
    print(f"  path（總）                  : {c['path']}"
          f"　（有 marker-end 箭頭頭：{c['path_arrow']}，其他（如箭頭標記定義本身）：{c['path_other']}）")
    print(f"  line（總）                  : {c['line']}"
          f"　（有 marker-end 箭頭頭：{c['line_arrow']}，裝飾線（分隔線／引導線）：{c['line_decorative']}）")
    print(f"  text（總）                  : {c['text']}"
          f"　（class=hypo 假說標籤：{c['text_hypo']}，一般文字：{c['text_plain']}）")
    print(f"  箭頭類小計（有 marker-end 的 path+line）: {c['arrow_like_total']}")
    print(f"  裝飾線小計（無箭頭頭，不計入箭頭類）    : {c['decorative_line_total']}")
    print(f"  文字類小計                  : {c['text_total']}")
    print(f"  元素總數                    : {c['total']}")


def _print_pptx_report(f: str, c: dict) -> None:
    print(f"\n[PPTX] {f}")
    print(f"  connector/freeform（箭頭類）: {c['connector_freeform']}")
    print(f"  textbox                     : {c['textbox']}")
    print(f"  autoshape 含文字            : {c['autoshape_with_text']}")
    print(f"  autoshape 無文字            : {c['autoshape_no_text']}")
    print(f"  其他                        : {c['other']}")
    print(f"  文字類小計（textbox+含文字autoshape）: {c['text_total']}")
    print(f"  段落總數（輔助參考，非主計數）: {c['paragraphs_total']}")
    print(f"  shape 總數                  : {c['shapes_total']}")


def main() -> int:
    ap = argparse.ArgumentParser(
        description="SVG／PPTX 元素清點工具——逐格式清點元素數，供回歸比對與新版式驗收使用。",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""範例：
  python count_elements.py fig.svg
  python count_elements.py fig.pptx
  python count_elements.py fig.svg fig.pptx --expect-svg 42 --expect-pptx 18
""",
    )
    ap.add_argument("files", nargs="+", help="要清點的 .svg 與／或 .pptx 檔案路徑（可一次給多個）")
    ap.add_argument("--expect-svg", type=int, default=None, help="斷言 SVG 元素總數（rect+path+line+text）")
    ap.add_argument("--expect-pptx", type=int, default=None, help="斷言 PPTX shape 總數")
    a = ap.parse_args()

    exit_code = 0
    svg_total = None
    pptx_total = None

    for f in a.files:
        low = f.lower()
        if low.endswith(".svg"):
            try:
                c = count_svg(f)
            except FileNotFoundError:
                print(f"錯誤：找不到檔案 {f}", file=sys.stderr)
                return 1
            except ET.ParseError as e:
                print(f"錯誤：SVG 解析失敗 {f}（{e}）", file=sys.stderr)
                return 1
            svg_total = c["total"]
            _print_svg_report(f, c)
        elif low.endswith(".pptx"):
            try:
                c = count_pptx(f)
            except ImportError:
                print("錯誤：清點 PPTX 需要 python-pptx（pip install python-pptx）", file=sys.stderr)
                return 1
            except FileNotFoundError:
                print(f"錯誤：找不到檔案 {f}", file=sys.stderr)
                return 1
            except Exception as e:  # noqa: BLE001 — 清點工具，任何讀檔異常都要誠實回報而非吞掉
                print(f"錯誤：讀取 PPTX 失敗 {f}（{e}）", file=sys.stderr)
                return 1
            pptx_total = c["shapes_total"]
            _print_pptx_report(f, c)
        else:
            print(f"錯誤：不支援的副檔名 {f}（僅支援 .svg / .pptx）", file=sys.stderr)
            return 1

    if svg_total is not None and pptx_total is not None:
        print("\n[對照] SVG 元素總數 vs PPTX shape 總數（兩者不需相等，僅供對照；"
              "務必分別檢查上方箭頭類／文字類是否兩邊都非零）")
        print(f"  SVG  元素總數 = {svg_total}")
        print(f"  PPTX shape 總數 = {pptx_total}")

    if a.expect_svg is not None:
        if svg_total is None:
            print("錯誤：--expect-svg 需要同時給一個 .svg 檔", file=sys.stderr)
            exit_code = 1
        elif svg_total != a.expect_svg:
            print(
                f"斷言失敗：SVG 元素總數 {svg_total} != 期望值 {a.expect_svg}"
                f"（差 {svg_total - a.expect_svg:+d}）",
                file=sys.stderr,
            )
            exit_code = 1

    if a.expect_pptx is not None:
        if pptx_total is None:
            print("錯誤：--expect-pptx 需要同時給一個 .pptx 檔", file=sys.stderr)
            exit_code = 1
        elif pptx_total != a.expect_pptx:
            print(
                f"斷言失敗：PPTX shape 總數 {pptx_total} != 期望值 {a.expect_pptx}"
                f"（差 {pptx_total - a.expect_pptx:+d}）",
                file=sys.stderr,
            )
            exit_code = 1

    return exit_code


if __name__ == "__main__":
    sys.exit(main())
